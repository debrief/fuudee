"""Synthetic Word and PDF documents for the narrative dialects' tests (DEV-05).

Never legacy Debrief's documents or their text: made-up, unclassified entries for
made-up platforms (BLUE ONE, RED TWO), organisations by initial only. The .docx are
built with python-docx (a development dependency only); the PDFs by hand, a line of
text a line, which is all the reader needs; the .doc files in tests/data were made
from these .docx by LibreOffice (see tests/data/README.md).
"""

from __future__ import annotations

import io
from collections.abc import Iterable

from docx import Document as WordDocument


def docx(
    paragraphs: Iterable[str | tuple[str, str, str]], table: list[list[str]] | None = None
) -> bytes:
    """A .docx of these paragraphs. A paragraph given as (shown, hidden, shown) has its
    middle hidden, as a narrative document hides its metadata after the DTG. A text's
    "\\n" is a soft line break; "\\t" a tab. A table, if given, comes first."""
    document = WordDocument()
    if table is not None:
        grid = document.add_table(rows=0, cols=max(len(r) for r in table))
        for row in table:
            cells = grid.add_row().cells
            if len(row) == 1:  # a heading across the row
                merged = cells[0].merge(cells[-1])
                merged.text = row[0]
                continue
            for cell, text in zip(cells, row, strict=False):
                _write(cell.paragraphs[0], text)
    for paragraph in paragraphs:
        p = document.add_paragraph()
        if isinstance(paragraph, tuple):
            shown, hidden, after = paragraph
            p.add_run(shown)
            p.add_run(hidden).font.hidden = True
            p.add_run(after)
        else:
            _write(p, paragraph)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def _write(paragraph, text: str) -> None:
    lines = text.split("\n")
    for i, line in enumerate(lines):
        run = paragraph.add_run()
        parts = line.split("\t")
        for j, part in enumerate(parts):
            run.add_text(part)
            if j < len(parts) - 1:
                run.add_tab()
        if i < len(lines) - 1:
            run.add_break()


def pdf(lines: list[str], per_page: int = 50) -> bytes:
    """A PDF with these lines of text, in Helvetica, so many to a page."""
    pages = [lines[i : i + per_page] for i in range(0, len(lines), per_page)] or [[]]
    objects: list[bytes] = []
    page_ids = [4 + 2 * i for i in range(len(pages))]
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{i} 0 R" for i in page_ids)
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    objects.append(
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>"
    )
    for i, page in enumerate(pages):
        content = ["BT", "/F1 10 Tf", "12 TL", "40 800 Td"]
        for line in page:
            escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            content.append(f"({escaped}) Tj T*")
        content.append("ET")
        stream = "\n".join(content).encode("cp1252")
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
            b"/Resources << /Font << /F1 3 0 R >> >> "
            + f"/Contents {page_ids[i] + 1} 0 R >>".encode()
        )
        objects.append(
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
        )
    return _pdf_file(objects)


def _pdf_file(objects: list[bytes]) -> bytes:
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(out)


def looping_pdf() -> bytes:
    """A PDF whose page tree lists itself among its pages: a reader must not recurse
    for ever (legacy Debrief's PDFBox once did)."""
    return _pdf_file(
        [
            b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [2 0 R 3 0 R] /Count 2 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] >>",
        ]
    )


# --- The narratives: made up ----------------------------------------------------------

# A narrative document with hidden metadata: what a .docx or .doc holds. The cases of
# legacy Debrief's tests: full entries with six- and four-figure DTGs, a late entry, a
# mangled day, a missing comma after a long type, short entries following on, a
# six-figure short entry into the next month, continuation text, a day marker.
NARRATIVE = [
    "Exercise B narrative, serial 2",
    "Prepared by the B. Squadron analysis cell",
    ("121005", ",12,03,20,HMS BLUE ONE,OOW COMMENT,", "Came to 090, 12 kn."),
    ("1106", ",12,03,2020,HMS BLUE ONE,COMMS,", "Signal from C. received."),
    ("122345", ",13,03,20,HMS BLUE ONE,OOW COMMENT,", "Late entry: contact faded."),
    ("110900", ",11,03,20,HMS BLUE ONE,OOW COMMENT,", "Mangled day in a cut and paste."),
    "Text continuing the mangled entry.",
    ("010800", ",01,04,20,RED TWO,LOG,", "New month for RED TWO."),
    ("020900", ",02,04,20,RED TWO,AVERYLONGTYPEWITHOUT", " COMMA and more text,trailing"),
    "1015 follow on text for RED TWO",
    "011030 SR023 B-123 R-5.1kyds C-321 S-6kts Classified Merchant vessel",
    "02 May 20",
    "1200 after the day marker",
    "",
]


def narrative_lines() -> list[str]:
    """The narrative document's paragraphs, as their full text (hidden parts shown)."""
    return ["".join(p) if isinstance(p, tuple) else p for p in NARRATIVE]


# The same log printed to PDF: the hidden metadata is not printed, so entries carry
# their DTG only, a day marker gives the date, and wrapped lines continue an entry.
NARRATIVE_PDF = [
    "Exercise B narrative, serial 2",
    "12 Mar 20",
    "121005 Came to 090, 12 kn.",
    "121106 Signal from C. received, read back",
    "and acknowledged.",
    "122345 Late entry: contact faded.",
    "130010 SR023 B-123 R-5.1kyds C-321 S-6kts Classified Merchant vessel",
    "130100 Exercise paused.",
]

# A Rider's narrative: a table, DTG Start first.
RIDER_TABLE = [
    ["DTG\nStart", "DTG Stop", "Serial", "Area", "Platform", "Remarks"],
    ["220000Z JUL 21", "222359Z JUL 21", "2", "B", "HMS Blue One", "-"],
    ["Chronological Narrative"],
    ["04:10:12", "273", "No Bm", "Contact gained."],
    ["04:11:13", "No B", "No Bm", "Contact held."],
    ["04:12:14", "274", "12", "Classified as merchant."],
    ["04:13:15", "", "21", "Beam only."],
    ["04:14:16", "92/112", "12/13", "Ambiguous bearing."],
    ["04:15:17", "096-112", "", "An arc."],
    ["", "", "", "No time given."],
    ["04:16:18", "", "", "First line\nSecond line"],
    ["23:59:00", "", "", "Before midnight."],
    ["00:01:00", "", "", "After midnight."],
    ["Footnote across the row"],
]

# ASW data: positions of one platform.
ASW = [
    "ASW DATA",
    "Prepared for exercise B",
    "TIMPD/261200ZAPR/IN/262359ZAPR/IN/APR/2021/APR/2021/",
    "TMPOS/261200ZAPR/IN/00.01.0N-000.21.0W/057T/04KTS/01.0M//",
    "TMPOS/261202ZAPR/IN/00.02.0N-000.22.0W/057T/04KTS/02.0M//",
    "TMPOS/261204ZAPR/IN/00.03.0N-000.23.0W/057T/04KTS/02.0M//",
    "TMPOS/261a206ZAPR/IN/00.05.0N-000.25.0W/057T/04KTS/01.0M//",
    "TMPOS/261208ZAPR/IN/00.07.0N-000.28.0W/057T/04KTS/01.0M//",
    "TMPOS/261212ZAPR/INS/00.10.30N-000.34.30W/180T/13KTS/000FT",
]

# Dated narrative lines.
DATED = [
    "Dated log",
    "2021,07,22,04:10:12,COMMENT,BLUE ONE,Contact gained.",
    "2021,07,22,04:11,CO_COMMENT,BLUE ONE,Closing, slowly.",
    "a stray line",
]

# RED ONE's log for the playground's sample, serial 3 (15 June 2023, 17:00 to 23:00):
# made up, as the sample is. BLUE TWO is a platform the plot has no track for; the last
# entry falls after the serial, so the reader is asked about both.
PLAYGROUND_LOG = [
    "Exercise B, serial 3: RED ONE's log",
    ("151705", ",15,06,23,RED ONE,OOW COMMENT,", "Dived. Course 270, 6 kn, 20 m."),
    (
        "151722",
        ",15,06,23,RED ONE,SONAR,",
        "Contact on the bow, bearing 095. Classified as BLUE ONE.",
    ),
    "Contact held on the towed array.",
    ("151740", ",15,06,23,RED ONE,OOW COMMENT,", "Came to 180 to open the contact."),
    ("151815", ",15,06,23,RED ONE,FCS,", "SR023 B-095 R-6kyds C-180 S-12kts Classified Frigate."),
    ("151830", ",15,06,23,BLUE TWO,COMMS,", "Signal from C. received: serial continues."),
    ("151902", ",15,06,23,RED ONE,SONAR,", "Contact lost astern."),
    ("160130", ",16,06,23,RED ONE,OOW COMMENT,", "Surfaced after the serial."),
]
