"""Legacy Debrief's Word and PDF narrative dialects, ported with their tests' cases
(CAP-19, EXT-04), on synthetic documents and lines only (DEV-05)."""

from pathlib import Path

import pytest
from synthetic_docs import (
    ASW,
    DATED,
    NARRATIVE,
    NARRATIVE_PDF,
    RIDER_TABLE,
    docx,
    looping_pdf,
    narrative_lines,
    pdf,
)

from debrief_io.documents import Document, DocumentError, read_document, sniff
from debrief_io.word import (
    ASW as ASW_DIALECT,
)
from debrief_io.word import (
    DATED as DATED_DIALECT,
)
from debrief_io.word import (
    DTG_MISSING,
    RIDER,
    NarrativeDocumentParser,
    asw_degrees,
    asw_fix,
    asw_hemisphere,
    asw_location,
    asw_time,
    asw_year,
    clean,
    day_marker,
    fcs_classified,
    fcs_element,
    fcs_range_yd,
    fcs_source,
    fcs_track,
    index_of_start,
    is_asw,
    parse_fcs,
    read_dialect,
    read_narrative_document,
    remove_bad_chars,
    rider_start,
    rider_text,
    rider_time,
)
from debrief_io.word import (
    NARRATIVE as NARRATIVE_DIALECT,
)
from debrief_kernel.time import format_utc, parse_utc

DATA = Path(__file__).parent / "data"


def when(ms: int) -> str:
    return format_utc(ms)[:16].replace("T", " ")


# --- The narrative document: a line at a time (testNarrEntryParsing) -------------------


def test_each_style_of_line_is_read_as_legacy_debrief_read_it() -> None:
    parser = NarrativeDocumentParser()
    lines = [
        "121005,12,03,20,BLUE ONE,SOURCE,Some text here",  # six-figure DTG, metadata
        "1106,12,03,2020,BLUE ONE,COMMS,Next message",  # four-figure DTG, four-figure year
        "122345,13,03,20,BLUE ONE,SOURCE,Late entry",  # 23xx, hidden day one ahead
        "110900,11,03,20,BLUE ONE,SOURCE,Mangled date",  # day back, same month
        "010800,01,04,85,RED TWO,TYPE,Old year",  # a 1980s year, new month
        "020900,02,04,85,RED TWO,VERYLONGTYPEWITHOUT COMMA and more text,trailing",
        "1015 follow on text for the entry",  # time only: date and platform before
        "011030 C0101 Tgt B-123 continued FCS text",  # six figures, day back: next month
        "just some more words for the previous entry",  # text continuing
        "12 Mar 20",  # a new day
        "1200 after the day marker",
        "1234",  # too short to be a time: text
    ]
    out = []
    for line in lines:
        read = parser.line(line)
        if read.day_marker:
            out.append("day")
            continue
        stamp = when(read.time_ms) if read.time_ms is not None else "no-dtg"
        out.append(f"{stamp}|{read.platform}|{read.entry_type}|{read.text}|{read.appended}")
    assert out == [
        "2020-03-12 10:05|BLUE ONE|SOURCE|Some text here|False",
        "2020-03-12 11:06|BLUE ONE|COMMS|Next message|False",
        "2020-03-12 23:45|BLUE ONE|SOURCE|Late entry|False",
        "2020-03-12 09:00|BLUE ONE|SOURCE|Mangled date|False",
        "1985-04-01 08:00|RED TWO|TYPE|Old year|False",
        # Legacy Debrief cut the type's last letter (VERYLONGTYPEWITHOU): the word is kept.
        "1985-04-02 09:00|RED TWO|VERYLONGTYPEWITHOUT|COMMA and more text,trailing|False",
        "1985-04-02 10:15|RED TWO|N/A|follow on text for the entry|False",
        "1985-05-01 10:30|RED TWO|FCS|C0101 Tgt B-123 continued FCS text|False",
        "no-dtg|None|None|just some more words for the previous entry|True",
        "day",
        "2020-03-12 12:00|RED TWO|N/A|after the day marker|False",
        "no-dtg|None|None|1234|True",
    ]
    assert len(parser.warnings) == 1 and "day went back at 0900" in parser.warnings[0]


def test_a_full_entry_has_its_platform_type_and_text_trimmed() -> None:
    read = NarrativeDocumentParser().line(
        "160909,16,09,2016,HMS BLUE ONE, CAT COMMENT, SOME COMMENT "
    )
    assert when(read.time_ms) == "2016-09-16 09:09"
    assert (read.platform, read.entry_type, read.text) == (
        "HMS BLUE ONE",
        "CAT COMMENT",
        "SOME COMMENT",
    )


def test_a_short_entry_takes_the_date_and_platform_before_it() -> None:
    parser = NarrativeDocumentParser()
    parser.line("160909,16,09,2016,HMS BLUE ONE, CAT COMMENT, SOME COMMENT ")
    read = parser.line("161006\tSOME COMMENT 2 ")
    assert when(read.time_ms) == "2016-09-16 10:06"
    assert (read.platform, read.text, read.appended) == ("HMS BLUE ONE", "SOME COMMENT 2", False)
    text = parser.line("SOME COMMENT ")
    assert (text.time_ms, text.platform, text.appended) == (None, None, True)


def test_a_time_with_nothing_before_it_is_not_read() -> None:
    from debrief_io.word import _LineError

    with pytest.raises(_LineError, match="no date or platform"):
        NarrativeDocumentParser().line("120506 irrelevant content 1")


def test_a_day_marker_dates_the_short_entries_after_it() -> None:
    # testParseDateFCS: a day marker alone gives no entry; a time after it does.
    parser = NarrativeDocumentParser()
    assert parser.line("12 Nov 2018").day_marker
    read = parser.line("120506 irrelevant content 1")
    assert when(read.time_ms) == "2018-11-12 05:06" and read.text == "irrelevant content 1"


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("12 Mar 20", "2020-03-12 00:00"),
        ("31 Dec 1995", "1995-12-31 00:00"),
        ("01 Jan 96", "1996-01-01 00:00"),
        ("15 SEP 16", "2016-09-15 00:00"),
        ("3 March 2021 and some words", "2021-03-03 00:00"),
        ("12 Mxr 20", None),
        ("31 Feb 20", None),
        ("irrelevant preamble", None),
        ("1234", None),
    ],
)
def test_day_markers(text: str, expected: str | None) -> None:
    found = day_marker(text)
    assert (when(found) if found is not None else None) == expected


def test_a_six_figure_day_back_in_december_is_january_next_year() -> None:
    parser = NarrativeDocumentParser()
    parser.line("310504,31,12,95,BLUE ONE,FCS,first")
    read = parser.line("010505 SR023 SOURCE_A FCS B-123 R-5.1kyds")
    assert when(read.time_ms) == "1996-01-01 05:05"


def test_a_day_that_does_not_exist_is_not_read_rather_than_rolled_on() -> None:
    from debrief_io.word import _LineError

    # Legacy Debrief's lenient calendar read 31 April as 1 May.
    with pytest.raises(_LineError, match="not a time that exists"):
        NarrativeDocumentParser().line("311005,31,04,21,BLUE ONE,LOG,No such day")


def test_controls_from_word_are_spaces_and_soft_returns_line_breaks() -> None:
    assert remove_bad_chars("a\x0bb\x01c\x13d\x15e\x05f") == "a\nb c d e f"


def test_the_start_of_records_marker() -> None:
    lines = ["preamble", "more", "START OF RECORDS FOR EXERCISE B", "x"]
    assert index_of_start(lines) == 2
    assert index_of_start(["no", "marker"]) == 0
    assert index_of_start(["  START RECORDS"]) == 0


# --- The narrative document: a document at a time --------------------------------------


def test_a_narrative_document_with_hidden_metadata() -> None:
    read = read_dialect(read_document("log.docx", docx(NARRATIVE)))
    assert read.dialect == NARRATIVE_DIALECT
    got = [(when(e.time_ms), e.source, e.entry_type) for e in read.entries]
    assert got == [
        ("2020-03-12 10:05", "HMS BLUE ONE", "OOW COMMENT"),
        ("2020-03-12 11:06", "HMS BLUE ONE", "COMMS"),
        ("2020-03-12 23:45", "HMS BLUE ONE", "OOW COMMENT"),
        ("2020-03-12 09:00", "HMS BLUE ONE", "OOW COMMENT"),
        ("2020-04-01 08:00", "RED TWO", "LOG"),
        ("2020-04-02 09:00", "RED TWO", "AVERYLONGTYPEWITHOUT"),
        ("2020-04-02 10:15", "RED TWO", None),  # legacy N/A: no type
        ("2020-05-01 10:30", "RED TWO", "FCS"),
        ("2020-05-02 12:00", "RED TWO", None),
    ]
    # Text continuing an entry joins it on a new line.
    assert (
        read.entries[3].text
        == "Mangled day in a cut and paste.\nText continuing the mangled entry."
    )
    assert read.preamble == 2
    assert read.types()["OOW COMMENT"] == 3 and read.types()[""] == 2
    fcs = read.entries[7].fcs
    assert fcs is not None and fcs.complete and fcs.range_yd == pytest.approx(5100)


def test_the_same_document_as_a_word_97_file() -> None:
    from_docx = read_dialect(read_document("log.docx", docx(NARRATIVE)))
    from_doc = read_dialect(
        read_document("log.doc", (DATA / "synthetic-narrative.doc").read_bytes())
    )
    assert [(e.time_ms, e.source, e.entry_type, e.text) for e in from_doc.entries] == [
        (e.time_ms, e.source, e.entry_type, e.text) for e in from_docx.entries
    ]


def test_the_docx_reader_reads_what_python_docx_reads() -> None:
    import io

    import docx as python_docx

    data = docx(NARRATIVE)
    theirs = [p.text for p in python_docx.Document(io.BytesIO(data)).paragraphs]
    assert read_document("log.docx", data).paragraphs == theirs
    assert theirs == narrative_lines()


def test_a_printed_log_has_no_metadata() -> None:
    read = read_dialect(read_document("log.pdf", pdf(NARRATIVE_PDF)))
    assert read.dialect == NARRATIVE_DIALECT
    assert [(when(e.time_ms), e.source, e.entry_type) for e in read.entries] == [
        ("2020-03-12 10:05", None, None),
        ("2020-03-12 11:06", None, None),
        ("2020-03-12 23:45", None, None),
        ("2020-03-13 00:10", None, "FCS"),
        ("2020-03-13 01:00", None, None),
    ]
    assert read.entries[1].text == "Signal from C. received, read back\nand acknowledged."


def test_continuation_keeps_every_line_up_to_the_next_entry() -> None:
    # Legacy Debrief joined six lines, dropped the seventh and joined on; product owner,
    # 15 Oct 2026: every line is kept, so a PDF's long entry (a line a paragraph) is whole.
    lines = ["121005,12,03,20,BLUE ONE,LOG,start"] + [f"more {i}" for i in range(40)]
    lines += ["121105,12,03,20,BLUE ONE,LOG,next"]
    read = read_narrative_document(lines)
    assert read.entries[0].text.count("\n") == 40
    assert read.entries[1].text == "next"
    assert read.overflow == 0


def test_continuation_stops_at_the_safety_cap_and_says_so(monkeypatch) -> None:
    import debrief_io.word as word

    monkeypatch.setattr(word, "MAX_CONTINUATION", 5)
    lines = ["121005,12,03,20,BLUE ONE,LOG,start"] + [f"more {i}" for i in range(8)]
    read = read_narrative_document(lines)
    assert read.entries[0].text.count("\n") == 5
    assert read.overflow == 3


# --- Out-of-order "times" are text (product owner, 15 Oct 2026) -------------------------


def entries_of(lines: list[str]) -> list[str]:
    """The document's entries as "dd HH:MM text", continuation lines joined by " / "."""
    read = read_narrative_document(lines)
    return [f"{when(e.time_ms)[8:]} {e.text.replace(chr(10), ' / ')}" for e in read.entries]


FIRST = "121400,12,03,20,BLUE ONE,OOW COMMENT,Contact bearing 090."


def test_a_number_out_of_order_ahead_is_text_of_the_entry_before() -> None:
    # Legacy Debrief read "2000 yds, closing." as an entry at 20:00, after 14:05.
    lines = [FIRST, "2000 yds, closing.", "1405 Came to 180.", "1410 Contact lost."]
    assert entries_of(lines) == [
        "12 14:00 Contact bearing 090. / 2000 yds, closing.",
        "12 14:05 Came to 180.",
        "12 14:10 Contact lost.",
    ]
    assert read_narrative_document(lines).as_text == [(2, "2000 yds, closing.")]


def test_numbers_in_a_row_do_not_vouch_for_each_other() -> None:
    lines = [FIRST, "2000 yds, closing.", "2100 yds, opening.", "1405 Came to 180.", "1410 Steady."]
    assert entries_of(lines) == [
        "12 14:00 Contact bearing 090. / 2000 yds, closing. / 2100 yds, opening.",
        "12 14:05 Came to 180.",
        "12 14:10 Steady.",
    ]


def test_a_number_out_of_order_behind_is_text_of_the_entry_before() -> None:
    # 05:00 is nine hours before 14:00: no late entry is that late.
    lines = [FIRST, "0500 yds, opening.", "1405 Came to 180."]
    assert entries_of(lines) == [
        "12 14:00 Contact bearing 090. / 0500 yds, opening.",
        "12 14:05 Came to 180.",
    ]
    # Thirteen hours before 18:00 is past midnight, on the 13th: then far ahead of 18:05.
    evening = ["121800,12,03,20,BLUE ONE,LOG,Dusk.", "0500 yds, opening.", "1805 Darkened."]
    assert entries_of(evening) == ["12 18:00 Dusk. / 0500 yds, opening.", "12 18:05 Darkened."]


def test_a_late_entry_within_the_hour_is_an_entry() -> None:
    lines = [FIRST, "1430 Came to 270.", "1415 Late entry: contact faded.", "1440 Steady."]
    assert entries_of(lines) == [
        "12 14:00 Contact bearing 090.",
        "12 14:30 Came to 270.",  # the late entry after it does not count against it
        "12 14:15 Late entry: contact faded.",
        "12 14:40 Steady.",
    ]
    # More than an hour apart, one of the two is out of order: the entry ahead of the
    # entry after it is text, as a number would be.
    later = [FIRST, "1530 Came to 270.", "1415 Late entry, written late.", "1540 Steady."]
    read = read_narrative_document(later)
    assert entries_of(later) == [
        "12 14:00 Contact bearing 090. / 1530 Came to 270.",
        "12 14:15 Late entry, written late.",
        "12 15:40 Steady.",
    ]
    assert read.as_text == [(2, "1530 Came to 270.")]


def test_midnight_rolls_over_without_a_day_marker() -> None:
    # Four figures: legacy Debrief read 0005 on the 12th, a day back; here past midnight.
    lines = [
        "122340,12,03,20,BLUE ONE,LOG,Late.",
        "2350 Later.",
        "0005 Past midnight.",
        "0015 Still on.",
    ]
    read = read_narrative_document(lines)
    assert entries_of(lines) == [
        "12 23:40 Late.",
        "12 23:50 Later.",
        "13 00:05 Past midnight.",
        "13 00:15 Still on.",
    ]
    assert read.as_text == []
    assert read.warnings == [
        "0005 is more than twelve hours before the entry before it: read as past midnight, "
        "on 13/03/2020."
    ]
    # Six figures carry the day; a day marker gives it.
    assert entries_of(
        ["122340,12,03,20,BLUE ONE,LOG,Late.", "130005 Past midnight.", "130015 Still on."]
    ) == ["12 23:40 Late.", "13 00:05 Past midnight.", "13 00:15 Still on."]
    assert entries_of(
        ["122340,12,03,20,BLUE ONE,LOG,Late.", "13 Mar 20", "0005 Past midnight.", "0015 Still on."]
    ) == ["12 23:40 Late.", "13 00:05 Past midnight.", "13 00:15 Still on."]


def test_a_late_entry_at_23xx_just_after_midnight_is_the_evening_before() -> None:
    # As legacy Debrief corrected a full entry's hidden day (correctLateEntryDay); in a
    # log with no metadata, a six-figure DTG's day back was the next month.
    for late in ("122345 Late entry: contact faded.", "2345 Late entry: contact faded."):
        lines = ["130010,13,03,20,BLUE ONE,LOG,Past midnight.", late, "130020 On."]
        assert entries_of(lines) == [
            "13 00:10 Past midnight.",
            "12 23:45 Late entry: contact faded.",
            "13 00:20 On.",
        ]


def test_full_entries_are_never_text_however_out_of_order() -> None:
    # The log goes on after an entry out of order: the entries after it agree.
    lines = [
        FIRST,
        "122000,12,03,20,BLUE ONE,LOG,Out of order.",
        "1405 Came to 180.",
        "1410 Steady.",
        "1415 Contact regained.",
    ]
    assert entries_of(lines) == [
        "12 14:00 Contact bearing 090.",
        "12 20:00 Out of order.",
        "12 14:05 Came to 180.",
        "12 14:10 Steady.",
        "12 14:15 Contact regained.",
    ]


def test_a_long_run_of_numbers_does_not_take_the_rest_of_the_log_with_it() -> None:
    # Three in a row are text: the entry after them is in order with the one before.
    three = [FIRST, "2000 yds, 091.", "2100 yds, 092.", "2200 yds, 093.", "1405 Came to 180."]
    assert entries_of([*three, "1410 Steady."]) == [
        "12 14:00 Contact bearing 090. / 2000 yds, 091. / 2100 yds, 092. / 2200 yds, 093.",
        "12 14:05 Came to 180.",
        "12 14:10 Steady.",
    ]
    # Four vouch for each other, and are entries, as legacy Debrief read them; the
    # entries after them agree with each other, not with the run, so they are entries.
    four = [*three[:4], "2300 yds, 094.", *three[4:], "1410 Steady.", "1415 Regained."]
    got = entries_of(four)
    assert got[-3:] == ["12 14:05 Came to 180.", "12 14:10 Steady.", "12 14:15 Regained."]
    assert len(got) == 8


def test_a_number_on_the_last_line_is_checked_against_the_entry_before_only() -> None:
    # Nothing after it to be out of order with: an entry, as legacy Debrief read it.
    assert entries_of([FIRST, "2000 yds, closing."]) == [
        "12 14:00 Contact bearing 090.",
        "12 20:00 yds, closing.",
    ]


def test_end_of_records_stops_reading_and_start_of_records_starts_it() -> None:
    lines = [
        "preamble 121005,12,03,20,X,Y,not an entry",
        "121005,12,03,20,BLUE ONE,LOG,before the start",
        "START OF RECORDS FOR EXERCISE B",
        "121006,12,03,20,BLUE ONE,LOG,first",
        "121007,12,03,20,BLUE ONE,LOG,second",
        "END OF RECORDS FOR EXERCISE B",
        "121008,12,03,20,BLUE ONE,LOG,after the end",
    ]
    read = read_narrative_document(lines)
    assert [e.text for e in read.entries] == ["first", "second"]
    assert read.stopped_at == "END OF RECORDS FOR EXERCISE B"


def test_a_bad_line_is_listed_and_reading_goes_on() -> None:
    read = read_narrative_document(
        [
            "121005,12,03,20,BLUE ONE,LOG,first",
            "121006,1x,03,20,BLUE ONE,LOG,bad",
            "121007,12,03,20,BLUE ONE,LOG,third",
        ]
    )
    assert [e.text for e in read.entries] == ["first", "third"]
    assert read.problems == ["Line 2 could not be read: its day, month and year are not numbers."]


# --- FCS entries (testParseFCS, testParseFCSRange, testAdvancedParseFCS) --------------

FCS_1 = (
    "160504,16,08,2016,BLUE ONE,FCS,   SR023 GGGG GGG (HHHH) B-123 R-5kyds C-321 S-6kts XX. "
    "Classified ABCD EFGH."
)
FCS_2 = "160403,16,09,2016,BLUE ONE,FCS, M01 1234 Rge B-311° R-12600 yds. Classified IJKL MNOP."
FCS_3 = (
    "160403,16,09,2016,BLUE ONE,FCS, M02 1234 Rge B-311° R-12.4kyds. Classified IJKL. "
    "Source from S333."
)
FCS_4 = (
    "160403,16,09,2016,BLUE ONE,FCS, M02 1234 Rge R-12.4kyds. Classified IJKL. Source from S333."
)


def _text(line: str) -> str:
    return NarrativeDocumentParser().line(line).text  # type: ignore[return-value]


def test_fcs_elements() -> None:
    assert fcs_element("B-", FCS_1) == 123
    assert fcs_element("C-", FCS_1) == 321
    assert fcs_element("R-", FCS_1) == 5
    assert fcs_element("S-", FCS_1) == 6
    assert fcs_element("B-", FCS_2) == 311
    assert fcs_element("R-", FCS_2) == 12600
    assert fcs_classified(FCS_1) == "ABCD EFGH."


def test_fcs_solutions() -> None:
    one = parse_fcs(_text(FCS_1))
    assert (one.range_yd, one.bearing_deg, one.contact, one.course_deg, one.speed_kn) == (
        5000,
        123,
        "023",
        321,
        6,
    )
    assert one.classification == "ABCD EFGH."
    two = parse_fcs(_text(FCS_2))
    assert (two.range_yd, two.bearing_deg, two.contact, two.course_deg, two.speed_kn) == (
        12600,
        311,
        "M01",
        0,
        0,
    )
    assert parse_fcs(_text(FCS_3)).contact == "M02"  # the master track before another number
    assert parse_fcs(_text(FCS_4)).bearing_deg is None
    assert not parse_fcs(_text(FCS_4)).complete


@pytest.mark.parametrize(
    ("text", "yards"),
    [
        ("SR023 B-123 R-5.1kyds C-321", 5100),
        ("SR023 B-123 R-5kyds C-321", 5000),
        ("SR023 B-123 R-800yds C-321 Classified X \r Y.", 800),
        ("SR023 B-123 R-800 m C-321 Classified X \nY.", 800 / 0.9144),
        ("M01 1234 Rge B-311° R-12.4kyds. Classified X.", 12400),
        ("M01 B-311 R-12 furlongs here", None),
    ],
)
def test_fcs_ranges(text: str, yards: float | None) -> None:
    got = fcs_range_yd(text)
    assert got == (pytest.approx(yards) if yards is not None else None)


@pytest.mark.parametrize(
    ("text", "track"),
    [
        ("asdfads S000 adf ag a", "000"),
        ("asdfads SV000 adf ag a", "000"),
        ("asdfads M00 adf ag a", "M00"),
        ("M00 0000", "M00"),
        ("asdfads adf ag a", None),
    ],
)
def test_fcs_track_numbers(text: str, track: str | None) -> None:
    assert fcs_track(text) == track


def test_fcs_sources() -> None:
    assert fcs_source("   SR023 SOURCE_A FCS B-123 R-5.1kyds C-321 S-6kts") == "SOURCE_A FCS"
    assert fcs_source("SR023 1936 GAINED FCS (AAAA) B-123 R-5kyds") == "1936"
    assert fcs_source("M01 GGGG GGGG GGG (HHHH) B-173 R-3.7kyds") == "GGGG GGGG GGG (HHHH)"
    assert fcs_source("SR023 by triangulation B-1") == "Triangulation"
    assert fcs_source("no track here B-1") is None


# --- Rider's narrative (TestImportRider) ---------------------------------------------


def test_a_riders_narrative_is_known_by_its_first_cell() -> None:
    doc = read_document("rider.docx", docx([], RIDER_TABLE))
    read = read_dialect(doc)
    assert read.dialect == RIDER
    assert len(read.entries) == 10
    assert {e.source for e in read.entries} == {"HMS Blue One"}
    assert {e.entry_type for e in read.entries} == {"Rider"}


def test_rider_entries_as_legacy_debrief_wrote_them() -> None:
    read = read_dialect(read_document("rider.docx", docx([], RIDER_TABLE)))
    e = read.entries
    assert (when(e[0].time_ms), e[0].text, e[0].bearing) == (
        "2021-07-22 04:10",
        "[Brg:273] Contact gained.",
        273,
    )
    assert (e[1].text, e[1].bearing) == ("Contact held.", None)
    assert e[2].text == "[Brg:274, Beam:12] Classified as merchant."
    assert e[3].text == "[Beam:21] Beam only."
    assert e[4].text == "[Brg:92/112, Beam:12/13] Ambiguous bearing."
    assert e[5].text == "[Bearing represents arc (096-112). Not imported.]An arc."
    # No time: the one before, a millisecond later, and said.
    assert e[6].time_ms == e[5].time_ms + 1 and e[6].text == DTG_MISSING + "No time given."
    assert e[7].text == "First line\nSecond line"
    # Past midnight: the next day (legacy Debrief kept the first day).
    assert when(e[9].time_ms) == "2021-07-23 00:01"
    assert read.warnings == ["Row 13, at 00:01:00, passed midnight: read as the next day."]
    assert read.problems == ["Row 14 has 1 cell, not the four of an entry: not read."]


def test_the_same_riders_narrative_as_a_word_97_file() -> None:
    from_docx = read_dialect(read_document("rider.docx", docx([], RIDER_TABLE)))
    from_doc = read_dialect(read_document("rider.doc", (DATA / "synthetic-rider.doc").read_bytes()))
    assert from_doc.dialect == RIDER
    assert [(e.time_ms, e.text) for e in from_doc.entries] == [
        (e.time_ms, e.text) for e in from_docx.entries
    ]


def test_rider_times() -> None:
    assert format_utc(rider_start("090000Z JUL 17")) == "2017-07-09T00:00:00.000Z"
    # Minutes and seconds of the start are set to zero (testTidyDate).
    assert rider_start("220001Z JUL 09") == rider_start("220000Z JUL 09")
    assert rider_time("02:01:34") == (2 * 3600 + 60 + 34) * 1000
    with pytest.raises(ValueError):
        rider_time("25:00:00")
    with pytest.raises(ValueError):
        rider_start("22 JUL 09")


def test_rider_text() -> None:
    assert rider_text("t", None, None, "21") == "[Beam:21] t"
    assert rider_text("t", 274, None, None) == "[Brg:274] t"
    assert rider_text("t", 92, 112, "12/13") == "[Brg:92/112, Beam:12/13] t"
    assert rider_text("t", None, None, None) == "t"


def test_a_rider_row_with_a_bad_time_or_bearing_is_listed() -> None:
    table = [*RIDER_TABLE[:3], ["4:10", "", "", "bad time"], ["04:11:13", "abc", "", "bad bearing"]]
    read = read_dialect(read_document("rider.docx", docx([], table)))
    assert [e.text for e in read.entries] == ["bad bearing"]
    assert read.problems == [
        "Row 4 could not be read: '4:10' is not a time like 04:10:12.",
        "Row 5: 'abc' is not a bearing.",
    ]


# --- ASW data (ImportASWDataDocument's TestImportWord) -------------------------------


def test_asw_data_is_known_by_its_position_lines() -> None:
    assert not is_asw(["Some old duff information"] * 20)
    assert is_asw(["Some old duff information"] * 20 + ["TMPOS/Here we go!"] * 10)
    read = read_dialect(read_document("asw.pdf", pdf(ASW)), track="BLUE ONE", default_year=2026)
    assert read.dialect == ASW_DIALECT
    assert len(read.fixes) == 5 and {f.track for f in read.fixes} == {"BLUE ONE"}
    assert read.problems == ["Line 7 could not be read: Badly formatted date:261a206ZAPR."]
    assert format_utc(read.fixes[0].time_ms) == "2021-04-26T12:00:00.000Z"


def test_asw_clean() -> None:
    assert clean("00O0") == "0000"
    assert clean(" 0Oo0") == "0000"
    assert clean("0ao0 ") == "0A00"
    assert clean("1,234") == "1.234"


def test_asw_dates_and_years() -> None:
    assert format_utc(asw_time("261200ZAPR", 2026)) == "2026-04-26T12:00:00.000Z"
    assert format_utc(asw_time("292258ZNOV", 2026)) == "2026-11-29T22:58:00.000Z"
    assert asw_year("TIMPD/290000Z/300200Z/APR/2016/APR2016") == 2016
    assert asw_year("TIMPD//230001Z//232359Z//APR//2017//APR//2017//") == 2017
    assert asw_year("TIMPD/261200ZAPR/IN/262359ZAPR/IN/APR/2019/APR/2019/") == 2019
    assert asw_year("TMPOS/290000Z/300200Z/APR/2016/APR2016") is None
    with pytest.raises(ValueError):
        asw_year("TIMPD/290000Z/300200Z/APR/16/APR2016")


def test_asw_hemispheres_and_degrees() -> None:
    assert asw_hemisphere("00.00.0N") == 1 and asw_hemisphere("000.00.00S") == -1
    assert asw_hemisphere("dddmmW") == -1 and asw_hemisphere("000.00. 00W") == -1
    assert asw_degrees("12.30.45", False) == pytest.approx(12.5125)
    assert asw_degrees("11545", False) == pytest.approx(115.75)
    assert asw_degrees("04 .30. 45", False) == pytest.approx(4.5125)
    assert asw_degrees("O11.3o.45", False) == pytest.approx(11.5125)
    assert asw_degrees("12.30.50", True) == pytest.approx(asw_degrees("12.30.30", False))
    with pytest.raises(ValueError, match="dd.mm.mm"):
        asw_degrees("011.30.75", False)


def test_asw_locations() -> None:
    assert asw_location("12.23.34N-121.12.1E", "") == pytest.approx(
        (12 + 23 / 60 + 34 / 3600, 121 + 12 / 60 + 1 / 3600)
    )
    assert asw_location("1230N/01215W", "") == pytest.approx((12.5, -12.25))
    assert asw_location("22.33N-020.11W", "") == pytest.approx((22.55, -(20 + 11 / 60)))
    for bad, why in [
        ("22.33.44J/020.11.22W", "Not a valid hemisphere:J"),
        ("22.33.44N=020.11.22W", "separator"),
        ("04.02.01N/211.22.11W", "Longitude out of limits"),
        ("94.02.01N/211.22.11W", "Latitude out of limits"),
    ]:
        with pytest.raises(ValueError, match=why):
            asw_location(bad, "")


def test_asw_fixes() -> None:
    one = asw_fix("TMPOS/261200ZAPR/IN/12.23.34N-121.12.1E/o57T/04KTS/05.1M//", 2019, "T")
    assert (one.course_deg, one.speed_kn, one.depth_m) == (57, 4, 5.1)
    assert format_utc(one.time_ms) == "2019-04-26T12:00:00.000Z"
    two = asw_fix("TIMPOS/262327ZAPR/INS/13.14.15S-011.22.33W/000T/13KTS/0o1FT", 2016, "T")
    assert two.depth_m == pytest.approx(0.3048) and two.speed_kn == 13
    three = asw_fix("TMPOS/220000ZAPR/GPS/1230N/01215W", 1998, "T")
    assert (three.course_deg, three.speed_kn, three.depth_m) == (0, 0, 0)


# --- Dated narrative (ImportWordDummy) and choosing the dialect -----------------------


def test_dated_narrative_lines() -> None:
    read = read_dialect(read_document("dated.docx", docx(DATED)))
    assert read.dialect == DATED_DIALECT
    assert [(format_utc(e.time_ms), e.entry_type, e.source, e.text) for e in read.entries] == [
        ("2021-07-22T04:10:12.000Z", "COMMENT", "BLUE ONE", "Contact gained."),
        ("2021-07-22T04:11:00.000Z", "CO_COMMENT", "BLUE ONE", "Closing, slowly."),
    ]
    assert read.preamble == 1 and read.problems == ["Line 4 is not an entry: it was not read."]


def test_documents_are_known_by_their_first_bytes_not_their_names() -> None:
    assert sniff(docx(["x"])) == "docx"
    assert sniff((DATA / "synthetic-rider.doc").read_bytes()) == "doc"
    assert sniff(pdf(["x"])) == "pdf"
    assert sniff(b";NARRATIVE: 110621 093015 X text") is None
    # A .docx named .doc, as legacy Debrief met one, is read as what it is.
    assert read_document("wrongly-named.doc", docx([], RIDER_TABLE)).format == "docx"


def test_damaged_and_hostile_documents_are_refused_in_words() -> None:
    with pytest.raises(DocumentError, match="not a Word"):
        read_document("x.txt", b"hello")
    with pytest.raises(DocumentError, match="damaged"):
        read_document("x.docx", b"PK\x03\x04broken")
    with pytest.raises(DocumentError, match="damaged"):
        read_document("x.doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 600)
    import io
    import zipfile

    hostile = io.BytesIO()
    with zipfile.ZipFile(hostile, "w") as archive:
        archive.writestr(
            "word/document.xml",
            '<?xml version="1.0"?><!DOCTYPE d [<!ENTITY a "aaaa">]><d>&a;</d>',
        )
    with pytest.raises(DocumentError, match="document type declaration"):
        read_document("x.docx", hostile.getvalue())


def test_a_pdf_with_no_text_says_it_may_be_a_scan() -> None:
    with pytest.raises(DocumentError, match="no text to read"):
        read_document("scan.pdf", pdf([]))


def test_a_pdf_whose_page_tree_loops_returns() -> None:
    # Legacy Debrief's PDFBox once recursed until the stack overflowed here.
    try:
        doc = read_document("loop.pdf", looping_pdf())
    except DocumentError:
        return
    assert doc.paragraphs == []


# --- A log with no metadata, in a longer document (testStartOf…, testDTGZ…, testMMSI…) -


def no_metadata() -> list[str]:
    """A printed log: day markers, six-figure short entries, text between; made up."""
    return [
        "title page",
        "distribution list",
        "31 Dec 1995",
        "310504 SR023 SOURCE A FCS B-123 R-5.1kyds C-321 S-6kts. Classified ZZZ.",
        "contents page",
        "glossary page",
        "01 Jan 1996",
        "010505 SR023 SOURCE A FCS B-123 R-5.1kyds C-321 S-6kts. Classified ZZZ.",
        "010506 nothing of note",
        "010507 SR023 SOURCE B FCS B-123 R-5kyds C-321 S-6kts. Classified ZZZ.",
        "010508 nothing more of note",
        "010509 SR023 SOURCE B FCS B-123 R-800yds C-321 S-6kts. Classified \r ZZZ.",
        "010608",
        "annex A",
        "annex B",
    ]


def test_a_log_with_no_metadata_reads_its_six_entries() -> None:
    read = read_narrative_document(no_metadata())
    assert len(read.entries) == 6  # testStartOfNotPresent
    assert {e.source for e in read.entries} == {None}
    # The day after 31 December is 1 January of the next year.
    assert when(read.entries[1].time_ms) == "1996-01-01 05:05"


def test_reading_starts_at_start_of_records_and_stops_at_end_of_records() -> None:
    lines = no_metadata()
    lines[5] = "START OF RECORDS FOR EXERCISE B"
    assert len(read_narrative_document(lines).entries) == 5  # testStartOfPresent
    lines[9] = "END OF RECORDS FOR EXERCISE B"
    assert len(read_narrative_document(lines).entries) == 2  # testStartOfAndEndOfPresent


def test_a_dtg_in_the_preamble_is_not_a_day() -> None:
    # testDTGZ_in_preamble: with no day before it, the first entry cannot be dated.
    lines = no_metadata()
    lines[2] = "121321Z some preamble"
    read = read_narrative_document(lines)
    assert len(read.entries) == 5
    assert read.problems == [
        "Line 4 could not be read: it has a time but no date or platform before it."
    ]


def test_an_mmsi_is_not_a_time() -> None:
    # testMMSI_in_preamble: eight figures are text, joined to the entry before.
    lines = no_metadata()
    lines[2] = lines[9] = "12632144 some preamble"
    read = read_narrative_document(lines)
    assert len(read.entries) == 4
    assert read.entries[1].text.endswith("\n12632144 some preamble")


def test_without_its_day_marker_the_next_day_is_the_next_month() -> None:
    # testMissingDateMarkerInNoMetadata: 01 after 31 December is January, in a new year.
    lines = no_metadata()
    del lines[6]
    read = read_narrative_document(lines)
    first, last = parse_utc("1996-01-01T05:00:00.000Z"), parse_utc("1996-01-01T05:59:00.000Z")
    assert len([e for e in read.entries if first <= e.time_ms <= last]) == 5


def test_platforms_are_matched_to_tracks_after_legacy_prefixes() -> None:
    from debrief_io.narrative import track_for_platform

    tracks = {"Nelson": "track-nelson", "Iron Duck": "track-iron-duck"}
    assert track_for_platform("HMS Boat", tracks) is None
    assert track_for_platform("HMS Nelson", tracks) == "track-nelson"
    assert track_for_platform("Hms Nelson", tracks) == "track-nelson"
    assert track_for_platform("RNAS Nelson", tracks) == "track-nelson"
    assert track_for_platform("Hms Iron Duck", tracks) == "track-iron-duck"
    assert track_for_platform("nelson", tracks) == "track-nelson"


def test_asw_lines_that_cannot_be_read_are_listed_by_number() -> None:
    # testFewFailures: three of twelve position lines spoilt, the rest read.
    good = "TMPOS/2612{m:02d}ZAPR/IN/00.{m:02d}.0N-000.21.0W/057T/04KTS/01.0M//"
    lines = ["HEADER"] * 4 + [good.format(m=m) for m in range(12)]
    lines[8] = lines[8].replace("2612", "261a2")
    lines[10] = lines[10].replace("057T", "05b7T")
    lines[12] = lines[12].replace("04KTS", "04KbTS")
    read = read_dialect(Document("docx", paragraphs=lines), track="T", default_year=2021)
    assert len(read.fixes) == 9
    assert [p.split(":")[0] for p in read.problems] == [
        "Line 9 could not be read",
        "Line 11 could not be read",
        "Line 13 could not be read",
    ]
    assert read.warnings == ["No TIMPD line gives the year: 2021 was taken."]
